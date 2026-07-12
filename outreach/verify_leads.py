#!/usr/bin/env python3
"""
verify_leads.py
Scores each lead's email-to-business-name match and splits into:
  - leads_verified.xlsx  (high/medium confidence — ready to use)
  - leads_flagged.xlsx   (low confidence — review manually)

Usage:
  cd outreach
  python3 verify_leads.py --leads leads_ready.xlsx
"""
import argparse
import re
from pathlib import Path
import openpyxl
from openpyxl import Workbook


def tokenise(text: str) -> set:
    """Break a string into meaningful lowercase words, strip noise."""
    if not text:
        return set()
    text = text.lower()
    # Remove common noise words that appear in both names and emails but mean nothing
    noise = {
        "the","and","of","for","in","at","on","a","an","co","ltd","llc","inc",
        "corp","services","service","solutions","group","company","business",
        "professional","pro","plus","premier","elite","best","top","local",
        "general","home","house","mr","ms","mrs","dr","info","contact","hello",
        "admin","office","mail","email","enquiries","enquiry","support","sales",
    }
    words = re.findall(r"[a-z0-9]+", text)
    return {w for w in words if len(w) > 2 and w not in noise}


def score_email(business_name: str, email: str) -> tuple[str, str]:
    """
    Returns (confidence, reason).
    confidence: 'high' | 'medium' | 'low'
    """
    if not email or not business_name:
        return "low", "missing data"

    email = email.lower().strip()
    local, _, domain = email.partition("@")
    domain_root = domain.split(".")[0] if domain else ""
    is_gmail = "gmail.com" in domain or "yahoo.com" in domain or "hotmail.com" in domain or "outlook.com" in domain or "icloud.com" in domain

    biz_tokens = tokenise(business_name)

    if is_gmail:
        # For Gmail, check overlap between local part and business name
        local_tokens = tokenise(local)
        overlap = biz_tokens & local_tokens
        if overlap:
            return "high", f"gmail matches: {', '.join(overlap)}"
        # Partial match — e.g. business is "ABC Plumbing" and email is "abcplumb123"
        for token in biz_tokens:
            if len(token) >= 4 and token[:4] in local:
                return "medium", f"gmail partial match on '{token}'"
        # Check reverse — local part words appear in business name
        for token in local_tokens:
            if len(token) >= 4 and any(token in b or b in token for b in biz_tokens):
                return "medium", f"gmail loose match on '{token}'"
        return "low", "gmail local part shares no words with business name"
    else:
        # Custom domain — check overlap between domain root and business name
        domain_tokens = tokenise(domain_root)
        overlap = biz_tokens & domain_tokens
        if overlap:
            return "high", f"domain matches: {', '.join(overlap)}"
        # Partial match on domain root
        for token in biz_tokens:
            if len(token) >= 4 and token[:4] in domain_root:
                return "medium", f"domain partial match on '{token}'"
        for token in domain_tokens:
            if len(token) >= 4 and any(token in b or b in token for b in biz_tokens):
                return "medium", f"domain loose match on '{token}'"
        # Generic domains (info@, contact@ etc) on a custom domain — probably fine
        generic_locals = {"info","contact","hello","admin","office","enquiries","enquiry","support","sales","mail"}
        if local in generic_locals:
            return "medium", f"generic inbox ({local}@) on custom domain — likely correct"
        return "low", f"domain '{domain_root}' shares no words with '{business_name}'"


def read_leads(path: str) -> tuple[list, list]:
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(row):
            rows.append(dict(zip(headers, row)))
    return headers, rows


def write_excel(path: str, headers: list, rows: list, extra_cols: list):
    wb = Workbook()
    ws = wb.active
    ws.append(headers + extra_cols)
    for row in rows:
        ws.append([row.get(h, "") for h in headers] + [row.get(c, "") for c in extra_cols])
    wb.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--leads", default="leads_ready.xlsx")
    args = parser.parse_args()

    input_path = Path(args.leads)
    if not input_path.exists():
        print(f"❌ File not found: {input_path}")
        return

    print(f"📂 Reading {input_path}...")
    headers, rows = read_leads(str(input_path))

    verified, flagged = [], []
    counts = {"high": 0, "medium": 0, "low": 0}

    for row in rows:
        biz   = str(row.get("business_name") or "")
        email = str(row.get("email") or "")
        confidence, reason = score_email(biz, email)
        row["email_confidence"] = confidence
        row["match_reason"]     = reason
        counts[confidence] += 1
        if confidence == "low":
            flagged.append(row)
        else:
            verified.append(row)

    extra = ["email_confidence", "match_reason"]

    out_verified = input_path.parent / "leads_verified.xlsx"
    out_flagged  = input_path.parent / "leads_flagged.xlsx"

    write_excel(str(out_verified), headers, verified, extra)
    write_excel(str(out_flagged),  headers, flagged,  extra)

    total = len(rows)
    print(f"\n{'─'*52}")
    print(f"✅ Total leads processed : {total}")
    print(f"   High confidence       : {counts['high']}")
    print(f"   Medium confidence     : {counts['medium']}")
    print(f"   Low confidence        : {counts['low']} → needs review")
    print(f"{'─'*52}")
    print(f"📄 leads_verified.xlsx   → {len(verified)} leads ready to use")
    print(f"🔍 leads_flagged.xlsx    → {len(flagged)} leads to review manually")
    print(f"\nNext steps:")
    print(f"  1. Open leads_flagged.xlsx — delete rows where email clearly doesn't match")
    print(f"  2. Run: python3 pipeline.py --leads leads_verified.xlsx --limit 3 --draft")
    print(f"  3. Check those 3 drafts in Gmail, then run the full batch")


if __name__ == "__main__":
    main()
