#!/usr/bin/env python3
"""
pipeline.py — Outreach Automation Pipeline
───────────────────────────────────────────
1. Reads leads from Excel
2. Upserts each lead into Supabase CRM
3. Generates a custom HTML mock site via Claude API
4. Screenshots with Puppeteer
5. Sends email from best available account (warm-up aware)
6. Updates CRM status automatically at every step

Usage:
  cd outreach
  export $(cat .env | xargs)
  python3 pipeline.py --leads leads_example.xlsx --limit 10

Flags:
  --limit N       Max leads to process today (default 10)
  --no-send       Generate sites + screenshots only, skip email
  --draft         Save as Gmail draft instead of sending
"""
import argparse
import json
import os
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl

from generate_site   import generate_mock_site
from gmail_draft     import send_email, create_draft, build_subject
from accounts        import pick_account, log_send, random_send_delay, is_business_hours, ACCOUNTS
from supabase_client import (upsert_lead, mark_site_generated, mark_email_sent, mark_draft_created,
                              increment_send_count, get_send_log, get_lead_by_email)

BASE  = Path(__file__).parent
SITES = BASE / "sites"
SHOTS = BASE / "screenshots"
SITES.mkdir(exist_ok=True)
SHOTS.mkdir(exist_ok=True)

YOUR_NAME     = os.environ.get("YOUR_NAME",     "Abdoul")
YOUR_WEBSITE  = os.environ.get("YOUR_WEBSITE",  "sitesbyabs.com")
COPY_VERSION  = os.environ.get("COPY_VERSION",  "v1")  # bump to v2, v3 to A/B test copy


def slug(name: str) -> str:
    return "".join(c if c.isalnum() else "_" for c in name.lower()).strip("_")


def read_leads(xlsx_path: str, limit: int) -> list[dict]:
    wb      = openpyxl.load_workbook(xlsx_path)
    ws      = wb.active
    headers = [c.value for c in ws[1]]
    leads   = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        lead = dict(zip(headers, row))
        if not lead.get("email") or not lead.get("business_name"):
            continue
        leads.append(lead)
        if len(leads) >= limit:
            break
    return leads


def take_screenshot(html_path: Path, png_path: Path) -> bool:
    result = subprocess.run(
        ["node", str(BASE / "screenshot.js"), str(html_path), str(png_path)],
        capture_output=True, text=True, cwd=str(BASE)
    )
    return result.returncode == 0


def run(leads_file: str, limit: int, no_send: bool = False, draft_mode: bool = False):
    leads    = read_leads(leads_file, limit)
    send_log = get_send_log()

    if not leads:
        print("No leads found.")
        return

    print(f"\n🚀 Processing {len(leads)} leads\n{'─'*52}")

    for i, lead in enumerate(leads, 1):
        name  = lead.get("business_name", f"Lead {i}")
        email = lead.get("email", "")
        key   = slug(name)

        print(f"\n[{i}/{len(leads)}] {name} <{email}>")

        # Normalise field names (Excel columns may use title case)
        if not lead.get("address"):
            lead["address"] = lead.get("Address", "")
        if not lead.get("pitch_type"):
            lead["pitch_type"] = "new"

        # ── 1. Upsert into Supabase ───────────────────────────────────────────
        lead_id = upsert_lead(lead)
        if not lead_id:
            # Already exists — get existing ID
            existing = get_lead_by_email(email)
            if existing:
                lead_id = existing["id"]
                already_sent    = existing["status"] == "email_sent"
                already_drafted = existing.get("email_sent_from") and existing["status"] != "email_sent"
                if already_sent or already_drafted:
                    print(f"  ✓ Already {'emailed' if already_sent else 'drafted'} — skipping")
                    continue
        print(f"  ✓ CRM: lead saved (id: {lead_id[:8]}...)")

        # ── 2. Generate mock site ─────────────────────────────────────────────
        html_file = SITES / f"{key}.html"
        if html_file.exists():
            print(f"  ✓ Site already generated — reusing")
        else:
            print(f"  → Generating mock site with Claude...")
            try:
                html = generate_mock_site(lead)
                html_file.write_text(html, encoding="utf-8")
                print(f"  ✓ Site generated ({len(html):,} chars)")
            except Exception as e:
                print(f"  ✗ Site generation failed: {e}")
                continue

        # ── 3. Screenshot ─────────────────────────────────────────────────────
        png_file = SHOTS / f"{key}.png"
        if not png_file.exists():
            print(f"  → Screenshotting...")
            if not take_screenshot(html_file, png_file):
                print(f"  ✗ Screenshot failed")
                continue
        print(f"  ✓ Screenshot ready")

        # Update CRM: site generated
        mark_site_generated(lead_id, str(png_file))

        if no_send:
            print(f"  ⏭ --no-send: skipping email")
            continue

        # ── 4. Pick sending account ───────────────────────────────────────────
        if draft_mode:
            # Drafts aren't real sends — round-robin accounts, ignore daily warm-up cap
            account = {**ACCOUNTS[(i - 1) % len(ACCOUNTS)], "remaining": "∞"}
        else:
            account = pick_account(send_log)
            if not account:
                print(f"  ⚠ All accounts hit daily limit — stopping")
                break
        print(f"  → Sending from: {account['email']} ({account['remaining']} left today)")

        # ── 5. Send / draft ───────────────────────────────────────────────────
        pitch   = lead.get("pitch_type", "new")
        subject = build_subject(lead)
        print(f"  → Pitch: {pitch} | Copy: {COPY_VERSION} | Subject: {subject}")
        try:
            if draft_mode:
                msg_id = create_draft(lead, str(png_file), account["name"], account["email"],
                                      YOUR_WEBSITE, COPY_VERSION, account=account)
                print(f"  ✓ Draft created")
            else:
                msg_id = send_email(lead, str(png_file), account, YOUR_WEBSITE, COPY_VERSION)
                print(f"  ✓ Email sent!")
        except Exception as e:
            print(f"  ✗ Email failed: {e}")
            continue

        # ── 6. Update CRM + send log ──────────────────────────────────────────
        if draft_mode:
            # Marks as "draft_created" (not "email_sent") so you flip it manually once you hit Send,
            # and so a rerun won't create a duplicate draft for this lead.
            mark_draft_created(lead_id, account["email"], subject, msg_id, COPY_VERSION, pitch)
            print(f"  ✓ CRM updated → Draft Created")
        else:
            mark_email_sent(lead_id, account["email"], subject, msg_id, COPY_VERSION, pitch)
            send_log = log_send(send_log, account["email"])
            increment_send_count(account["email"])
            print(f"  ✓ CRM updated → Email Sent")

        # ── 7. Human delay before next send ──────────────────────────────────
        if i < len(leads) and not draft_mode:
            remaining = sum(a.get("remaining", 0) for a in [account])
            delay = random_send_delay(remaining)
            mins  = delay // 60
            secs  = delay % 60
            print(f"  ⏱ Next send in {mins}m {secs}s...")
            time.sleep(delay)

    # ── Summary ───────────────────────────────────────────────────────────────
    print(f"\n{'─'*52}")
    print(f"✅ Done! Check your CRM: brilliant-gnome-d72d80.netlify.app")
    print(f"📧 Emails sent today per account:")
    for acct in ACCOUNTS:
        today = datetime.now().strftime("%Y-%m-%d")
        count = send_log.get(acct["email"], {}).get(today, 0)
        print(f"   {acct['email']}: {count} sent")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--leads",    default="leads_ready.xlsx")
    parser.add_argument("--limit",    type=int, default=10)
    parser.add_argument("--no-send",  action="store_true")
    parser.add_argument("--draft",    action="store_true")
    args = parser.parse_args()

    missing = [k for k in ["ANTHROPIC_API_KEY", "SUPABASE_URL", "SUPABASE_KEY"] if not os.environ.get(k)]
    if missing:
        print(f"❌ Missing env vars: {', '.join(missing)}")
        print("   Run: export $(cat .env | xargs)")
        sys.exit(1)

    run(args.leads, args.limit, no_send=args.no_send, draft_mode=args.draft)
