#!/usr/bin/env python3
"""
audit_coverage.py
Reports how many verified leads have real photos vs gradient fallback,
and breaks down by industry.

Usage:
  cd outreach
  python3 audit_coverage.py --leads leads_verified.xlsx
"""
import argparse
from collections import Counter
from pathlib import Path
import openpyxl
from hero_images import find_real_photo

TARGET_KEYWORDS = [
    "plumb", "roof", "gutter", "chimney",
    "electr", "eletric",
    "clean", "maid", "janitor", "upholstery",
    "lands", "lawn", "tree", "arborist", "sod",
    "hvac", "heat", "air condition", "mechanical",
    "handy", "remodel", "construct", "contractor", "fence", "concrete", "demo",
    "paint",
    "pest", "exterminator",
    "floor",
    "water damage", "restoration",
    "detail", "car wash",
]

def is_target_niche(industry):
    low = (industry or "").lower()
    return any(kw in low for kw in TARGET_KEYWORDS)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--leads", default="leads_verified.xlsx")
    args = parser.parse_args()

    path = Path(args.leads)
    if not path.exists():
        print(f"❌ File not found: {path}")
        return

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    rows = [dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(r)]

    total = len(rows)
    skipped = []
    has_photo = []
    has_gradient = []
    industry_counter = Counter()

    for row in rows:
        industry = str(row.get("industry") or "")
        if not is_target_niche(industry):
            skipped.append(industry)
            continue
        photo = find_real_photo(industry)
        industry_counter[industry.strip()] += 1
        if photo:
            has_photo.append(industry)
        else:
            has_gradient.append(industry)

    print(f"\n{'─'*56}")
    print(f"📂 {path.name}: {total} leads")
    print(f"{'─'*56}")
    print(f"⏭  Skipped (not target niche):  {len(skipped)}")
    print(f"📸 Will have real photo:         {len(has_photo)}")
    print(f"🎨 Will use gradient fallback:   {len(has_gradient)}")
    print(f"{'─'*56}")

    if has_gradient:
        gradient_industries = Counter(has_gradient)
        print(f"\n🎨 Gradient industries (consider adding photos):")
        for ind, count in gradient_industries.most_common():
            print(f"   {count:>4}x  {ind}")

    if skipped:
        skipped_counter = Counter(skipped)
        print(f"\n⏭  Skipped industries (not in target niches):")
        for ind, count in skipped_counter.most_common(15):
            print(f"   {count:>4}x  {ind or '(blank)'}")

    print(f"\n✅ Leads that will be processed: {len(has_photo) + len(has_gradient)}")
    print(f"   → {len(has_photo)} with real hero photo")
    print(f"   → {len(has_gradient)} with gradient background (still looks professional)")

if __name__ == "__main__":
    main()
