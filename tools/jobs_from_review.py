"""Turn the client's edited HCC-Jobs-Review.xlsx back into a CSV for Job Board → Import jobs (CSV).

Usage: python3 tools/jobs_from_review.py client-docs/HCC-Jobs-Review.xlsx client-docs/HCC-Jobs-Import.csv out.csv
The original import CSV supplies each advert's formatted (HTML) description; if the client rewrote the
description text in the workbook, their text is used instead. Rows with no job ref are new jobs.
"""
import csv, datetime, re, sys, html
from openpyxl import load_workbook

XLSX, ORIG, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
orig = {r['ref']: r for r in csv.DictReader(open(ORIG, encoding='utf-8-sig'))}
ws = load_workbook(XLSX, data_only=True)['Jobs']
head = [str(c.value or '').strip() for c in ws[1]]
col = lambda name: head.index(name)
H = {'status': 'Live or Closed?', 'title': 'Job title', 'location': 'Location (town or city)', 'pay': 'Pay shown on the site', 'pay_min': 'Lowest hourly rate (number only)',
     'shifts': 'Shift pattern (e.g. Days, Nights)', 'contract': 'Contract', 'sector': 'Sector', 'job_type': 'Job type', 'closing': 'Closing date (optional)',
     'badge': 'Badge (optional)', 'ref': 'Job ref (do not change)', 'slug': 'Web address (do not change)', 'description': 'Description (optional to edit)', 'notes': 'Notes for us'}
plain = lambda s: re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', ' ', s or ''))).strip()
out, notes = [], []
for row in ws.iter_rows(min_row=2, values_only=True):
    g = lambda k: row[col(H[k])] if col(H[k]) < len(row) else None
    title = str(g('title') or '').strip()
    if not title:
        continue
    ref = str(g('ref') or '').strip()
    desc_x = str(g('description') or '').strip()
    o = orig.get(ref, {})
    desc = o.get('description', '')
    px = plain(desc_x)
    if px and not plain(desc).startswith(px[:max(0, len(px) - 10)]):  # client rewrote it (the workbook holds a shortened plain-text copy)
        desc = '\n\n'.join('<p>' + html.escape(p.strip()) + '</p>' for p in re.split(r'\n\s*\n', desc_x) if p.strip())
    closing = g('closing')
    closing = closing.strftime('%Y-%m-%d') if isinstance(closing, (datetime.date, datetime.datetime)) else str(closing or '').strip()
    pmin = g('pay_min')
    out.append([ref, title, g('sector') or '', g('job_type') or '', g('location') or '', g('contract') or 'Temporary', g('pay') or '',
                '' if pmin in (None, '') else pmin, g('shifts') or '', str(g('badge') or '').lower(), closing,
                'live' if str(g('status') or '').strip().lower() == 'live' else 'expired', (str(g('slug') or '').strip() if 'slug' in H and H['slug'] in head else o.get('slug', '')), desc])
    if g('notes'):
        notes.append('%s: %s' % (ref or title, g('notes')))
with open(OUT, 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(['ref', 'title', 'sector', 'job_type', 'location', 'contract', 'pay', 'pay_min', 'shifts', 'badge', 'closing', 'status', 'slug', 'description'])
    w.writerows(out)
print(len(out), 'jobs written,', sum(r[11] == 'live' for r in out), 'live.')
if notes:
    print('Client notes:\n  ' + '\n  '.join(notes))
