"""Build the client review workbook and an importer-ready CSV from the current site's public job feed."""
import json, re, html, csv, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# Usage: python3 tools/jobs_review.py <folder with full.json and all_*.json from the old site REST API> <out.xlsx> <out.csv>
SRC, OUT_X, OUT_C = sys.argv[1], sys.argv[2], sys.argv[3]
J = json.load(open(SRC + '/full.json'))
terms = {t: {x['id']: x for x in json.load(open(f'{SRC}/all_{t}.json'))} for t in ['location', 'job_type', 'job_sector', 'contract']}
NLW = 12.71  # UK National Living Wage from April 2026 (age 21+)
TODAY = datetime.date(2026, 10, 28)
FIX_TYPE = {'Senior Care Assiatant': 'Senior Care Assistant', 'Sinior Care Assistant': 'Senior Care Assistant', 'Health Care Assistant': 'Healthcare Assistant'}

def txt(s): return html.unescape(re.sub(r'<[^>]+>', ' ', s or '')).replace('\xa0', ' ')
def clean_html(s):
    s = re.sub(r'<(script|style)[^>]*>.*?</\1>', '', s or '', flags=re.S | re.I)
    s = re.sub(r'<(/?)(\w+)[^>]*>', lambda m: '<%s%s>' % (m.group(1), m.group(2).lower()) if m.group(2).lower() in ('p', 'h2', 'h3', 'h4', 'ul', 'ol', 'li', 'strong', 'b', 'em', 'br') else '', s)
    s = s.replace('&nbsp;', ' ').replace('\xa0', ' ')
    s = re.sub(r'<p>\s*</p>', '', s)
    return re.sub(r'\n{3,}', '\n\n', s).strip()

def names(job, tax):
    ids = job.get(tax, [])
    t = terms[tax]
    if tax == 'location':  # most specific places only (drop a parent when its child is also tagged)
        parents = {t[i].get('parent') for i in ids if i in t}
        ids = [i for i in ids if i not in parents]
    return [html.unescape(t[i]['name']) for i in ids if i in t]

rows = []
for j in sorted(J, key=lambda x: x['date'], reverse=True):
    title = html.unescape(j['title']['rendered'])
    body = txt(j['content']['rendered'])
    m = re.search(r'£\s?([\d.,]+)(?:\s?(?:–|-|to)\s?£?\s?([\d.,]+))?\s*(per hour|p/h|ph|an hour|/hr|per annum|a year)?', body)
    pay, pmin = '', ''
    if m:
        lo = m.group(1).rstrip('.,'); hi = (m.group(2) or '').rstrip('.,'); unit = m.group(3) or 'per hour'
        pay = '£%s%s %s' % (lo, ' – £' + hi if hi else '', 'per hour' if unit in ('p/h', 'ph', 'an hour', '/hr') else unit)
        try: pmin = round(float(lo.replace(',', '')), 2)
        except ValueError: pmin = ''
    types = [FIX_TYPE.get(x, x) for x in names(j, 'job_type')]
    sectors = names(j, 'job_sector')
    sector = 'Healthcare' if 'Healthcare' in sectors else 'Chefs' if ('Chef Jobs' in sectors or 'Chef' in types) else 'Housekeeping & Cleaning' if 'Housekeeping' in types else 'Hospitality'
    contracts = names(j, 'contract')
    contract = 'Permanent' if contracts == ['Permanent'] else 'Temp to perm' if 'Temporary to Permanent' in contracts else 'Temporary'
    posted = datetime.date.fromisoformat(j['date'][:10])
    age = (TODAY - posted).days
    flags = []
    if age > 183: flags.append('Posted %s (%d months ago): still open?' % (posted.strftime('%b %Y'), age // 30))
    if isinstance(pmin, float) and pmin < NLW and 'annum' not in pay and 'year' not in pay: flags.append('Pay below minimum wage (£%.2f from April 2026): update the rate' % NLW)
    if not pay: flags.append('No pay found in the advert: please add')
    loc = names(j, 'location')
    if not loc:  # newest adverts: the town is in the title, e.g. "... Jobs in Maryport, CA15, Cumbria"
        g = re.search(r'^(?:.*\bin\s+)?([A-Z][A-Za-z&\s]+?)(?:\s*[,(]|\s*$)', title.split(' – ')[-1].strip())
        if g and not re.search(r'Agency|Work|Jobs|Temp', g.group(1)):
            loc = [g.group(1).strip()]
            flags.append('Location taken from the title: please check')
        else:
            flags.append('No location: please add')
    rows.append({
        'status': 'Live' if age <= 183 else 'Closed', 'title': title, 'location': ' | '.join(loc), 'pay': pay, 'pay_min': pmin, 'shifts': '',
        'contract': contract, 'sector': sector, 'job_type': ' | '.join(dict.fromkeys(types)), 'closing': '', 'badge': '', 'client_notes': '',
        'flags': '\n'.join(flags), 'posted': posted.strftime('%d/%m/%Y'), 'link': j['link'], 'slug': j['slug'], 'ref': 'HCC-W%d' % j['id'], 'description': clean_html(j['content']['rendered']),
    })

# ---- Workbook for the client ----
wb = Workbook(); ws = wb.active; ws.title = 'Jobs'
cols = [  # key, header, width, editable
    ('status', 'Live or Closed?', 14, True), ('title', 'Job title', 46, True), ('location', 'Location (town or city)', 24, True),
    ('pay', 'Pay shown on the site', 26, True), ('pay_min', 'Lowest hourly rate (number only)', 16, True), ('shifts', 'Shift pattern (e.g. Days, Nights)', 22, True),
    ('contract', 'Contract', 15, True), ('sector', 'Sector', 22, True), ('job_type', 'Job type', 22, True), ('closing', 'Closing date (optional)', 15, True),
    ('badge', 'Badge (optional)', 12, True), ('client_notes', 'Notes for us', 30, True), ('flags', 'Please check (from us)', 44, False),
    ('posted', 'First posted', 12, False), ('link', 'Current advert', 40, False), ('ref', 'Job ref (do not change)', 16, False), ('slug', 'Web address (do not change)', 30, False), ('description', 'Description (optional to edit)', 60, True),
]
edit_fill = PatternFill('solid', fgColor='FFFFFF'); lock_fill = PatternFill('solid', fgColor='EEF3F2'); head_fill = PatternFill('solid', fgColor='171A45')
flag_fill = PatternFill('solid', fgColor='FFF2CC'); thin = Side(style='thin', color='D5DEDC')
for c, (k, h, w, ed) in enumerate(cols, 1):
    cell = ws.cell(row=1, column=c, value=h); cell.font = Font(bold=True, color='FFFFFF'); cell.fill = head_fill
    cell.alignment = Alignment(wrap_text=True, vertical='center'); ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[1].height = 36
for r, row in enumerate(rows, 2):
    for c, (k, h, w, ed) in enumerate(cols, 1):
        v = row[k]
        if k == 'description': v = txt(v).strip()[:3000]  # readable text; the HTML version stays in the import file
        cell = ws.cell(row=r, column=c, value=v)
        cell.alignment = Alignment(wrap_text=k in ('title', 'flags', 'description', 'client_notes'), vertical='top')
        cell.fill = flag_fill if (k == 'flags' and v) else (edit_fill if ed else lock_fill)
        cell.border = Border(bottom=thin)
        if k == 'link': cell.hyperlink = v; cell.font = Font(color='2D338C', underline='single')
        if k == 'pay_min' and v != '': cell.number_format = '£0.00'
    ws.row_dimensions[r].height = 60
ws.freeze_panes = 'C2'; ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cols)), len(rows) + 1)
n = len(rows) + 1
for col, opts in (('A', '"Live,Closed"'), ('G', '"Temporary,Permanent,Temp to perm"'), ('H', '"Healthcare,Chefs,Hospitality,Housekeeping & Cleaning,Admin & Reception"'), ('K', '"New,Urgent"')):
    dv = DataValidation(type='list', formula1=opts, allow_blank=True); dv.add('%s2:%s%d' % (col, col, n)); ws.add_data_validation(dv)
dvn = DataValidation(type='decimal', operator='between', formula1='0', formula2='500', allow_blank=True, error='Number only, e.g. 13.50'); dvn.add('E2:E%d' % n); ws.add_data_validation(dvn)
dvd = DataValidation(type='date', operator='greaterThan', formula1='DATE(2026,1,1)', allow_blank=True, error='Enter a date, e.g. 31/12/2026'); dvd.add('J2:J%d' % n); ws.add_data_validation(dvd)
for r in range(2, n + 1): ws.cell(row=r, column=10).number_format = 'DD/MM/YYYY'

info = wb.create_sheet('How to use', 0)
lines = [
    ('HCC-Temps job adverts: review before they go on the new website', True),
    ('', False),
    ('These are the %d job adverts on the current website. Please check each row on the "Jobs" tab and send the file back.' % len(rows), False),
    ('', False),
    ('1. Live or Closed? Choose Live for jobs still open. Choose Closed for jobs that are filled or no longer needed (they will not show on the website).', False),
    ('2. Pay: update "Pay shown on the site" (e.g. £13.50 – £14.00 per hour) and "Lowest hourly rate" (number only, e.g. 13.50).', False),
    ('3. Location: the town or city. For more than one, separate with |  e.g. Leeds | Bradford', False),
    ('4. Shift pattern, closing date and badge (New / Urgent) are optional.', False),
    ('5. The yellow "Please check" column shows what we spotted: old adverts, missing pay, or pay below the minimum wage (£%.2f an hour from April 2026).' % NLW, False),
    ('6. To add a new job, add a row at the bottom and leave "Job ref" empty.', False),
    ('7. Grey columns are for reference. Please don\'t change "Job ref" or "Web address": they link each row to its advert.', False),
    ('', False),
    ('We have suggested Live for adverts posted in the last 6 months and Closed for older ones. Please change these to what is right today.', False),
    ('Summary: %d adverts · %d suggested Live · %d with pay below minimum wage · %d with no pay found.' % (len(rows), sum(r['status'] == 'Live' for r in rows), sum('minimum wage' in r['flags'] for r in rows), sum('No pay' in r['flags'] for r in rows)), True),
]
info.column_dimensions['A'].width = 120
for i, (l, b) in enumerate(lines, 1):
    c = info.cell(row=i, column=1, value=l); c.font = Font(bold=b, size=14 if i == 1 else 11, color='171A45' if b else '1E2140'); c.alignment = Alignment(wrap_text=True)
wb.active = 0
wb.save(OUT_X)

# ---- Importer-ready CSV (same columns as Job Board → Import jobs) ----
with open(OUT_C, 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(['ref', 'title', 'sector', 'job_type', 'location', 'contract', 'pay', 'pay_min', 'shifts', 'badge', 'closing', 'status', 'slug', 'description'])
    for r in rows:
        w.writerow([r['ref'], r['title'], r['sector'], r['job_type'], r['location'], r['contract'], r['pay'], r['pay_min'], '', '', '', 'live' if r['status'] == 'Live' else 'expired', r['slug'], r['description']])
# Same jobs, all live: for trying the job board before the client's review comes back.
with open(OUT_C.replace('.csv', '-AllLive.csv'), 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(['ref', 'title', 'sector', 'job_type', 'location', 'contract', 'pay', 'pay_min', 'shifts', 'badge', 'closing', 'status', 'slug', 'description'])
    for r in rows:
        w.writerow([r['ref'], r['title'], r['sector'], r['job_type'], r['location'], r['contract'], r['pay'], r['pay_min'], '', '', '', 'live', r['slug'], r['description']])
print(len(rows), 'jobs |', sum(r['status'] == 'Live' for r in rows), 'live |', sum('minimum wage' in r['flags'] for r in rows), 'below NLW |', sum('No pay' in r['flags'] for r in rows), 'no pay |', sum('No location' in r['flags'] for r in rows), 'no location')
